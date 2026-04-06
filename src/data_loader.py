"""
DataLoader — 文档加载模块

支持格式：PDF / TXT / MD / DOCX / HTML

改进（策略一 + 策略五）：
- 策略一：提取 PDF Outline，建立页码 → 标题路径映射，供切分器注入 section_path
- 策略五：过滤目录页，避免目录条目干扰检索
"""

import os
import re
import uuid
from pathlib import Path
from typing import List, Optional, Dict, Any, Tuple
from dataclasses import dataclass, field
from loguru import logger


# =========================================
# 文档数据结构
# =========================================
@dataclass
class Document:
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    doc_id: str = field(default_factory=lambda: str(uuid.uuid4()))

    def __repr__(self):
        preview = self.content[:60] + "..." if len(self.content) > 60 else self.content
        return f"Document(id={self.doc_id[:8]}, chars={len(self.content)}, preview='{preview}')"


# =========================================
# 文档加载器
# =========================================
class DataLoader:

    SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".html", ".htm"}

    # 目录页关键词（策略五）
    _TOC_PATTERN = re.compile(
        r'^(contents|目录|目次|table\s+of\s+contents|致谢|acknowledgements?)$',
        re.IGNORECASE
    )

    def __init__(self, config: Optional[Dict] = None):
        self.config = config or {}
        self.max_file_size_mb = self.config.get("max_file_size", 50)

    def load_document(self, path) -> Document:
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        size_mb = path.stat().st_size / (1024 * 1024)
        if size_mb > self.max_file_size_mb:
            raise ValueError(f"File too large: {size_mb:.1f}MB (max {self.max_file_size_mb}MB)")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported format: {ext}")

        loaders = {
            ".pdf":  self._load_pdf,
            ".txt":  self._load_text,
            ".md":   self._load_text,
            ".docx": self._load_docx,
            ".html": self._load_html,
            ".htm":  self._load_html,
        }

        content, metadata = loaders[ext](path)
        metadata.update({
            "source": str(path),
            "filename": path.name,
            "extension": ext,
            "file_size_mb": round(size_mb, 3),
        })

        logger.info(f"Loaded: {path.name} ({len(content)} chars)")
        return Document(content=content, metadata=metadata)

    def load_directory(self, directory, recursive: bool = True) -> List[Document]:
        directory = Path(directory)
        if not directory.is_dir():
            raise NotADirectoryError(f"Not a directory: {directory}")

        pattern = "**/*" if recursive else "*"
        files = [
            f for f in directory.glob(pattern)
            if f.is_file() and f.suffix.lower() in self.SUPPORTED_EXTENSIONS
        ]
        logger.info(f"Found {len(files)} files in {directory}")

        documents = []
        for f in files:
            try:
                documents.append(self.load_document(f))
            except Exception as e:
                logger.warning(f"Skipping {f.name}: {e}")

        logger.info(f"Loaded {len(documents)}/{len(files)} documents")
        return documents

    # --------------------------------------------------
    # PDF 加载（含策略一 Outline 提取 + 策略五目录过滤）
    # --------------------------------------------------
    def _load_pdf(self, path: Path) -> Tuple[str, Dict]:
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("pypdf not installed. Run: pip install pypdf")

        reader = PdfReader(str(path))
        total_pages = len(reader.pages)

        # 策略一：提取书签目录，建立 0-based 页码 → 标题名称 映射
        outline_entries = self._extract_outline(reader)
        page_to_section = self._build_page_to_section(outline_entries, total_pages)
        if outline_entries:
            logger.info(f"Outline: {len(outline_entries)} entries from {path.name}")

        # 策略五：识别目录页范围，后续跳过
        toc_pages = self._detect_toc_pages(reader, total_pages)
        if toc_pages:
            logger.info(f"TOC pages filtered: {sorted(toc_pages)} in {path.name}")

        pages = []
        for i, page in enumerate(reader.pages):
            # 策略五：跳过目录页
            if i in toc_pages:
                continue

            text = page.extract_text() or ""
            if text.strip():
                text = self._clean_pdf_text(text)
                # 插入页码标记（1-based），切分器据此追踪页码
                pages.append(f"[PAGE:{i+1}] {text}")

        content = "\n\n".join(pages)
        metadata = {
            "page_count": total_pages,
            "type": "pdf",
            # page_to_section 传给切分器使用，不直接存入 chunk metadata
            "_page_to_section": page_to_section,
        }
        return content, metadata

    # --------------------------------------------------
    # 策略一：提取 PDF Outline
    # --------------------------------------------------
    def _extract_outline(self, reader) -> List[Dict]:
        """
        递归遍历 PDF 书签树，返回按页码排序的条目列表。
        每项：{'title': str, 'page': int, 'level': int}
        """
        entries = []

        def _walk(items, level=1):
            for item in items:
                if isinstance(item, list):
                    _walk(item, level + 1)
                else:
                    try:
                        page_num = reader.get_destination_page_number(item)
                        title = (item.title or "").strip()
                        if title:
                            entries.append({'title': title, 'page': page_num, 'level': level})
                    except Exception:
                        pass

        try:
            _walk(reader.outline)
        except Exception as e:
            logger.debug(f"Outline extraction failed: {e}")

        return sorted(entries, key=lambda x: x['page'])

    def _build_page_to_section(self, entries: List[Dict], total_pages: int) -> Dict[int, str]:
        """
        按页码区间建立 0-based page → section_title 映射。
        chunk 所在页落在哪两个相邻书签之间，就归属于前一个书签。
        """
        if not entries:
            return {}
        mapping = {}
        for i, entry in enumerate(entries):
            end = entries[i + 1]['page'] if i + 1 < len(entries) else total_pages
            for p in range(entry['page'], end):
                mapping[p] = entry['title']
        return mapping

    # --------------------------------------------------
    # 策略五：检测目录页
    # --------------------------------------------------
    def _detect_toc_pages(self, reader, total_pages: int) -> set:
        """
        识别目录页：找到含目录关键词的页面，并把其后连续的目录条目页也纳入过滤范围。
        只扫描前 20% 的页面，避免误伤正文。
        """
        toc_pages = set()
        scan_limit = max(1, int(total_pages * 0.2))

        for i in range(min(scan_limit, total_pages)):
            text = (reader.pages[i].extract_text() or "").strip()
            first_line = text.split('\n')[0].strip() if text else ""
            # 去除全角空格后匹配
            normalized = re.sub(r'[\s\u3000]+', '', first_line)
            if self._TOC_PATTERN.match(normalized):
                toc_pages.add(i)
                # 把紧随其后、内容像目录条目的页也过滤掉（最多连续 10 页）
                for j in range(i + 1, min(i + 11, total_pages)):
                    next_text = reader.pages[j].extract_text() or ""
                    # 目录条目特征：大量省略号或页码数字结尾
                    if re.search(r'\.{3,}|\s\d{1,4}\s*$', next_text, re.MULTILINE):
                        toc_pages.add(j)
                    else:
                        break
                break  # 一个文档通常只有一个目录

        return toc_pages

    # --------------------------------------------------
    # PDF 文本清洗
    # --------------------------------------------------
    def _clean_pdf_text(self, text: str) -> str:
        text = re.sub(r'(?<!\n)\n(?!\n)', ' ', text)
        text = re.sub(r'[ \t]+', ' ', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        text = re.sub(r'(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])', '', text)
        text = re.sub(r'(?<=[\u4e00-\u9fff])\s+(?=[，。！？、；：""''（）【】])', '', text)
        text = re.sub(r'(?<=[，。！？、；：""''（）【】])\s+(?=[\u4e00-\u9fff])', '', text)
        return text.strip()

    # --------------------------------------------------
    # TXT / MD
    # --------------------------------------------------
    def _load_text(self, path: Path) -> Tuple[str, Dict]:
        for enc in ["utf-8", "utf-8-sig", "gbk", "latin-1"]:
            try:
                content = path.read_text(encoding=enc)
                return content, {"type": path.suffix.lstrip(".")}
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Cannot decode file: {path}")

    # --------------------------------------------------
    # DOCX
    # --------------------------------------------------
    def _load_docx(self, path: Path) -> Tuple[str, Dict]:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError("python-docx not installed. Run: pip install python-docx")

        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n\n".join(paragraphs)
        return content, {"paragraph_count": len(paragraphs), "type": "docx"}

    # --------------------------------------------------
    # HTML
    # --------------------------------------------------
    def _load_html(self, path: Path) -> Tuple[str, Dict]:
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            raise ImportError("beautifulsoup4 not installed. Run: pip install beautifulsoup4")

        html = path.read_text(encoding="utf-8", errors="ignore")
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style", "nav", "footer"]):
            tag.decompose()
        content = soup.get_text(separator="\n", strip=True)
        title = soup.title.string if soup.title else ""
        return content, {"title": title, "type": "html"}
